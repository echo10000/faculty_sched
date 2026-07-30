from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from faculty.services import compute_department_load_summary


DAYS = ("MON", "TUE", "WED", "THU", "FRI", "SAT")


def export_block_timetable_pdf(block):
    """Return a printable single-block weekly timetable PDF as bytes."""
    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        leftMargin=0.45 * inch,
        rightMargin=0.45 * inch,
        topMargin=0.45 * inch,
        bottomMargin=0.45 * inch,
    )
    styles = getSampleStyleSheet()
    assignments = list(block.assignments.select_related("subject", "room").all())
    times = sorted({(assignment.start_time, assignment.end_time) for assignment in assignments})
    by_time_and_day = {
        (assignment.start_time, assignment.end_time, assignment.day_of_week): assignment
        for assignment in assignments
    }
    table_rows = [["Time", *DAYS]]
    if times:
        for start, end in times:
            row = [f"{start:%H:%M}-{end:%H:%M}"]
            for day in DAYS:
                assignment = by_time_and_day.get((start, end, day))
                row.append(
                    Paragraph(
                        f"<b>{assignment.subject.code}</b><br/>{assignment.room.name}"
                        if assignment else "",
                        styles["BodyText"],
                    )
                )
            table_rows.append(row)
    else:
        table_rows.append(["No assignments have been scheduled.", "", "", "", "", "", ""])

    table = Table(table_rows, colWidths=[1.05 * inch] + [1.45 * inch] * 6, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#BFC7D5")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BACKGROUND", (0, 1), (0, -1), colors.HexColor("#EAF0F6")),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))
    document.build([
        Paragraph(f"Weekly Timetable - {block}", styles["Title"]),
        Spacer(1, 0.18 * inch),
        table,
    ])
    return buffer.getvalue()


def export_faculty_load_report_xlsx(department, term):
    """Return a formatted department faculty-load workbook as bytes."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Faculty Load Report"
    headers = [
        "Faculty",
        "Base Load",
        "Designation Release",
        "Assigned Units",
        "Required Load",
        "Difference",
        "Status",
    ]
    worksheet.append(headers)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    for summary in compute_department_load_summary(department, term):
        worksheet.append([
            str(summary["faculty"]),
            summary["base_load"],
            summary["units_released"],
            summary["assigned_units"],
            summary["required_load"],
            summary["difference"],
            summary["status"],
        ])
    for column in range(2, 7):
        for cell in worksheet.iter_cols(min_col=column, max_col=column, min_row=2):
            cell[0].number_format = "0.0"
    for index, header in enumerate(headers, start=1):
        longest = max(len(str(cell.value or "")) for cell in worksheet[get_column_letter(index)])
        worksheet.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 14), 32)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
