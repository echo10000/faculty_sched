"""Source-agnostic renderers for a tabular report.

Each renderer accepts a mapping with ``title`` and ``source_label`` strings,
``meta`` as a list of ``(label, value)`` pairs, ``headers`` as a list of
strings, ``rows`` as a list of equally sized lists of scalar values,
``filename_base`` as a string, and ``orientation`` as either ``portrait`` or
``landscape``. The caller supplies an already scoped data set. ``filename_base``
is reserved for the caller's download filename; renderers do not use it.

``render_csv`` returns a UTF-8 CSV string containing only the column header and
data rows. ``render_xlsx`` and ``render_pdf`` return complete document bytes.
"""

import csv
from datetime import date, datetime, time
from io import BytesIO, StringIO
from numbers import Number
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle
from django.utils import timezone


_BLUE = "1F4E78"
_FORMULA_PREFIXES = ("=", "+", "-", "@")


def _as_text(value):
    return "" if value is None else str(value)


def _spreadsheet_safe(value):
    """Keep numbers numeric while disarming formulas in all text fields."""
    if isinstance(value, Number) and not isinstance(value, complex):
        return value
    if isinstance(value, datetime) and timezone.is_aware(value):
        return timezone.localtime(value).replace(tzinfo=None)
    if isinstance(value, time) and value.tzinfo is not None:
        return value.replace(tzinfo=None)
    if isinstance(value, (date, datetime, time)):
        return value
    text = _as_text(value)
    if text.lstrip().startswith(_FORMULA_PREFIXES):
        return "'" + text
    return text


def _pdf_paragraph(value, style):
    # Paragraph interprets XML-like markup, so all report content is escaped.
    return Paragraph(escape(_as_text(value)).replace("\n", "<br/>"), style)


def render_csv(report):
    """Return UTF-8 CSV text with the column header and data rows only.

    Text that a spreadsheet might execute as a formula is prefixed with an
    apostrophe. Numeric scalars, including negative numbers, retain their value.
    """
    output = StringIO(newline="")
    writer = csv.writer(output, lineterminator="\r\n")
    writer.writerow([_spreadsheet_safe(value) for value in report["headers"]])
    writer.writerows([_spreadsheet_safe(value) for value in row] for row in report["rows"])
    return output.getvalue()


def render_xlsx(report):
    """Return formatted XLSX bytes with title, source, metadata, and table."""
    workbook = Workbook()
    sheet = workbook.active
    forbidden = '[]:*?/\\'
    sheet_name = "".join(" " if char in forbidden else char for char in report["title"]).strip()
    sheet.title = sheet_name[:31].strip("'") or "Report"

    headers = report["headers"]
    column_count = max(1, len(headers))
    last_column = get_column_letter(column_count)
    if column_count > 1:
        sheet.merge_cells(f"A1:{last_column}1")
    sheet["A1"] = _spreadsheet_safe(report["title"])
    sheet["A1"].font = Font(size=16, bold=True, color=_BLUE)
    sheet.row_dimensions[1].height = 27
    if column_count > 1:
        sheet.merge_cells(f"A2:{last_column}2")
    sheet["A2"] = _spreadsheet_safe(report["source_label"])
    sheet["A2"].font = Font(size=11, italic=True, color="526579")

    for row_number, (label, value) in enumerate(report["meta"], start=3):
        if column_count > 1:
            sheet.cell(row_number, 1, _spreadsheet_safe(label)).font = Font(bold=True)
            sheet.cell(row_number, 2, _spreadsheet_safe(value))
        else:
            sheet.cell(row_number, 1, _spreadsheet_safe(f"{_as_text(label)}: {_as_text(value)}"))

    header_row = 4 + len(report["meta"])
    fill = PatternFill("solid", fgColor=_BLUE)
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(header_row, col, _spreadsheet_safe(header))
        cell.fill = fill
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[header_row].height = 30

    widths = [max(14, min(48, len(_as_text(header)) + 3)) for header in headers]
    for row_number, row in enumerate(report["rows"], start=header_row + 1):
        for col, value in enumerate(row, start=1):
            cell = sheet.cell(row_number, col, _spreadsheet_safe(value))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if col <= len(widths):
                widths[col - 1] = max(widths[col - 1], min(48, len(_as_text(value)) + 2))
    for col, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(col)].width = width

    sheet.freeze_panes = f"A{header_row + 1}"
    if headers:
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{max(header_row, sheet.max_row)}"
    sheet.sheet_view.showGridLines = False
    sheet.print_title_rows = f"1:{header_row}"
    sheet.page_setup.orientation = report["orientation"]
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def render_pdf(report):
    """Return printable A4 PDF bytes with repeating table headers and pages."""
    output = BytesIO()
    pagesize = landscape(A4) if report["orientation"] == "landscape" else A4
    document = SimpleDocTemplate(
        output,
        pagesize=pagesize,
        leftMargin=14 * mm,
        rightMargin=14 * mm,
        topMargin=20 * mm,
        bottomMargin=16 * mm,
        title=_as_text(report["title"]),
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("ReportTitle", parent=styles["Title"], fontSize=14, leading=17)
    source_style = ParagraphStyle("ReportSource", parent=styles["Normal"], fontSize=9, leading=12)
    meta_style = ParagraphStyle("ReportMeta", parent=styles["Normal"], fontSize=8, leading=11)
    cell_style = ParagraphStyle("ReportCell", parent=styles["Normal"], fontSize=8, leading=10)
    header_style = ParagraphStyle(
        "ReportHeader", parent=cell_style, textColor=colors.white, fontName="Helvetica-Bold"
    )

    story = [
        _pdf_paragraph(report["title"], title_style),
        _pdf_paragraph(report["source_label"], source_style),
        Spacer(1, 4 * mm),
    ]
    for label, value in report["meta"]:
        story.append(_pdf_paragraph(f"{_as_text(label)}: {_as_text(value)}", meta_style))
    if report["meta"]:
        story.append(Spacer(1, 5 * mm))

    headers = report["headers"]
    if headers:
        available = pagesize[0] - document.leftMargin - document.rightMargin
        # Distribute space by the observed content, while keeping narrow fields
        # usable. Paragraph cells wrap within the allocated width.
        samples = [headers, *report["rows"][:100]]
        weights = [
            max(8, min(32, max(len(_as_text(row[index])) for row in samples if index < len(row))))
            for index in range(len(headers))
        ]
        widths = [available * weight / sum(weights) for weight in weights]
        data = [[_pdf_paragraph(value, header_style) for value in headers]]
        data.extend([_pdf_paragraph(value, cell_style) for value in row] for row in report["rows"])
        table = LongTable(data, colWidths=widths, repeatRows=1, splitInRow=1, hAlign="LEFT")
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{_BLUE}")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F6FA")]),
            ("LINEBELOW", (0, 0), (-1, 0), 0.5, colors.HexColor(f"#{_BLUE}")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        story.append(table)

    def page_footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#526579"))
        canvas.drawString(doc.leftMargin, pagesize[1] - 12 * mm, _as_text(report["source_label"]))
        canvas.drawRightString(pagesize[0] - doc.rightMargin, 9 * mm, f"Page {doc.page}")
        canvas.restoreState()

    document.build(story, onFirstPage=page_footer, onLaterPages=page_footer)
    return output.getvalue()
